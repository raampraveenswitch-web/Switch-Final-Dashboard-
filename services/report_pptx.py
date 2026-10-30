"""PowerPoint export (16:9) in the SWITCH dashboard style. Charts are native PowerPoint charts, so numbers stay
editable in PowerPoint; the funnel is drawn with shapes. Built with python-pptx (no browser or Office needed)."""
from __future__ import annotations

import copy
import io
from datetime import datetime

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import nsdecls, qn
from pptx.oxml import parse_xml
from pptx.util import Emu, Inches, Pt

from config import APP_NAME, BRAND_BLUE, BRAND_GREEN, BRAND_TEAL, LOGO_PATH, PALETTE, TAGLINE
from services import presentation as P
from services.report_data import MIN_GROUP_NOTE, Chart, Report

W, H = 13.333, 7.5
FONT = "Segoe UI"
NAVY, GREEN, TEAL = BRAND_BLUE, BRAND_GREEN, BRAND_TEAL
INK, MUTED, LINE, BG, CARD = "#0F1F3D", "#5E6B85", "#DCE3EE", "#F3F6FA", "#FFFFFF"
GOOD, BAD = "#1E8E4E", "#C93C3C"
NUMFMT = {"int": "#,##0", "dec1": "#,##0.0", "dec2": "0.00", "pct": '0.0"%"'}
LOGO_RATIO = 753 / 142          # width / height of assets/switch_logo.png


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h.lstrip("#").upper())


def blend(c1: str, c2: str, t: float) -> str:
    a = [int(c1[i:i + 2], 16) for i in (1, 3, 5)]
    b = [int(c2[i:i + 2], 16) for i in (1, 3, 5)]
    return "#" + "".join(f"{round(x + (y - x) * t):02x}" for x, y in zip(a, b))


# ------------------------------------------------------------------ primitives
def _box(slide, x, y, w, h, fill=CARD, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06, lw=0.75):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius * min(1.0, 1.2 / max(min(w, h), 0.1))
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def _text(slide, x, y, w, h, text, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          italic=False, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.0)
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        runs = ln if isinstance(ln, list) else [(ln, {})]
        for t, st in runs:
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = FONT
            f.size = Pt(st.get("size", size))
            f.bold = st.get("bold", bold)
            f.italic = st.get("italic", italic)
            f.color.rgb = rgb(st.get("color", color))
    return tb


def _logo(slide, x, y, h):
    if LOGO_PATH.exists():
        slide.shapes.add_picture(str(LOGO_PATH), Inches(x), Inches(y), height=Inches(h))
    else:
        _text(slide, x, y, h * LOGO_RATIO, h, [[("SW", {"color": NAVY}), ("//", {"color": GREEN}), ("TCH", {"color": NAVY})]],
              size=h * 52, bold=True)


def _bg(slide, color=BG):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = rgb(color)


def _header(slide, rep: Report, page_title: str = ""):
    _bg(slide)
    band = _box(slide, 0, 0, W, 1.08, fill=CARD, line=None, shape=MSO_SHAPE.RECTANGLE)
    band.line.fill.background()
    _logo(slide, 0.45, 0.3, 0.46)
    x = 0.45 + 0.46 * LOGO_RATIO + 0.28
    bar = _box(slide, x, 0.26, 0.04, 0.56, fill=NAVY, line=None, shape=MSO_SHAPE.RECTANGLE)
    _box(slide, x, 0.54, 0.04, 0.28, fill=GREEN, line=None, shape=MSO_SHAPE.RECTANGLE)
    _text(slide, x + 0.22, 0.22, 6.2, 0.45, rep.title + (f"  ·  {page_title}" if page_title else ""), 22, NAVY, True)
    _text(slide, x + 0.22, 0.64, 6.2, 0.3, TAGLINE, 12, GREEN, True)
    chip = _box(slide, W - 3.75, 0.33, 3.3, 0.46, fill="#F3F6FA", line=LINE, radius=0.5)
    loc = rep.location if rep.location and rep.location != "All" else "All locations"
    _text(slide, W - 3.7, 0.33, 3.2, 0.46, f"{rep.period}  ·  {loc}", 11.5, INK, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    _box(slide, 0, 1.08, W, 0.035, fill=NAVY, line=None, shape=MSO_SHAPE.RECTANGLE)
    _box(slide, 0, 1.08, 2.2, 0.035, fill=GREEN, line=None, shape=MSO_SHAPE.RECTANGLE)
    return bar


def _footer(slide, rep: Report, n: int):
    _text(slide, 0.45, H - 0.38, 6, 0.25, f"{APP_NAME}  ·  Confidential", 9, MUTED)
    _text(slide, W / 2 - 3, H - 0.38, 6, 0.25, f"{rep.period_range}  ·  Source: {rep.source or 'workbook'}", 9, MUTED,
          align=PP_ALIGN.CENTER)
    _text(slide, W - 1.45, H - 0.38, 1.0, 0.25, str(n), 9, MUTED, align=PP_ALIGN.RIGHT)


def _card(slide, x, y, w, h, title, subtitle=""):
    _box(slide, x, y, w, h)
    _text(slide, x + 0.22, y + 0.16, w - 0.44, 0.32, title, 14, INK, True)
    if subtitle:
        _text(slide, x + 0.22, y + 0.47, w - 0.44, 0.26, subtitle, 10, MUTED)


# ------------------------------------------------------------------ chart styling helpers
def _font(obj, size=9, color=MUTED, bold=False):
    obj.font.size = Pt(size)
    obj.font.name = FONT
    obj.font.bold = bold
    obj.font.color.rgb = rgb(color)


def _axis_clean(ax, show_labels=True, grid=True):
    ax.has_major_gridlines = grid
    if grid:
        ax.major_gridlines.format.line.color.rgb = rgb("#E6EBF3")
        ax.major_gridlines.format.line.width = Pt(0.75)
    ax.format.line.fill.background()
    ax.tick_labels.font.size = Pt(9)
    ax.tick_labels.font.name = FONT
    ax.tick_labels.font.color.rgb = rgb(MUTED)
    from pptx.enum.chart import XL_TICK_MARK
    ax.major_tick_mark = XL_TICK_MARK.NONE
    ax.minor_tick_mark = XL_TICK_MARK.NONE
    if not show_labels:
        ax.visible = False


def _labels(plot, fmt="int", pos=XL_LABEL_POSITION.OUTSIDE_END, color=INK, size=9):
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format = NUMFMT.get(fmt, "#,##0")
    dl.number_format_is_linked = False
    dl.font.size = Pt(size)
    dl.font.name = FONT
    dl.font.color.rgb = rgb(color)
    try:
        dl.position = pos
    except Exception:
        pass
    dl.show_value = True


def _legend(chart, pos=XL_LEGEND_POSITION.TOP):
    chart.has_legend = True
    chart.legend.position = pos
    chart.legend.include_in_layout = False
    chart.legend.font.size = Pt(9)
    chart.legend.font.name = FONT
    chart.legend.font.color.rgb = rgb(INK)


def _int_axis(ax, series):
    """Whole-number ticks for small integer counts (avoids 0.2, 0.4 … on a count axis)."""
    vals = [v for _, vs in series for v in vs if v is not None]
    if vals and all(float(v).is_integer() for v in vals):
        ax.tick_labels.number_format = "#,##0"
        ax.tick_labels.number_format_is_linked = False
        mx = ax.maximum_scale or max(vals)
        if mx <= 12:
            ax.major_unit = 1 if mx <= 6 else 2


def _top(vals) -> float:
    v = [x for x in vals if x is not None]
    return max(v) if v and max(v) > 0 else 1.0


def _clean_vals(vals):
    return [None if v is None else round(v, 4) for v in vals]


def _add_chart(slide, kind, x, y, w, h, cats, series):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, _clean_vals(vals))
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), cd)
    ch = gf.chart
    ch.font.name = FONT
    ch.font.size = Pt(9)
    ch.has_title = False
    return ch


# ------------------------------------------------------------------ chart types
def _combo(slide, c: Chart, x, y, w, h):
    series = list(c.series) + ([c.line] if c.line else [])
    ch = _add_chart(slide, XL_CHART_TYPE.COLUMN_CLUSTERED, x, y, w, h, c.cats, series)
    plot = ch.plots[0]
    plot.gap_width = 70
    plot.overlap = -8
    cols = [NAVY, TEAL, "#5B8BD9", "#7FCFCB"]
    for i, s in enumerate(list(plot.series)[: len(c.series)]):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(cols[i % len(cols)])
        s.invert_if_negative = False
    bar_top = max(_top(v) for _, v in c.series) if c.series else 1.0
    _axis_clean(ch.value_axis)
    _axis_clean(ch.category_axis, grid=False)
    ch.value_axis.minimum_scale = 0
    if not c.line:
        _labels(plot, c.fmt)
        ch.value_axis.maximum_scale = bar_top * 1.2
        _int_axis(ch.value_axis, c.series)
    if len(series) > 1:
        _legend(ch)
    if not c.line:
        return ch
    # --- convert the last series into a labelled line (optionally on a secondary axis).
    # Primary-axis limits are set first: python-pptx's value_axis is ambiguous once a second valAx exists.
    ln_vals0 = [v for v in c.line[1] if v is not None]
    ch.value_axis.maximum_scale = bar_top * 2.1 if c.secondary else max(bar_top, max(ln_vals0 or [1.0])) * 1.28
    _int_axis(ch.value_axis, c.series)
    pa = ch._chartSpace.chart.plotArea
    bar = pa.find(qn("c:barChart"))
    sers = bar.findall(qn("c:ser"))
    src = sers[-1]
    bar.remove(src)
    n = len(sers) - 1
    ln_vals = [v for v in c.line[1] if v is not None]
    ln_top = max(ln_vals) if ln_vals else 1.0
    ln_lo = min(ln_vals) if ln_vals else 0.0
    color = GREEN.lstrip("#")
    ser = parse_xml(f'<c:ser {nsdecls("c", "a")}><c:idx val="{n}"/><c:order val="{n}"/></c:ser>')
    ser.append(copy.deepcopy(src.find(qn("c:tx"))))
    ser.append(parse_xml(f'<c:spPr {nsdecls("c", "a")}><a:ln w="28575" cap="rnd"><a:solidFill><a:srgbClr val="{color}"/>'
                         f'</a:solidFill><a:round/></a:ln></c:spPr>'))
    ser.append(parse_xml(f'<c:marker {nsdecls("c", "a")}><c:symbol val="circle"/><c:size val="6"/><c:spPr><a:solidFill>'
                         f'<a:srgbClr val="{color}"/></a:solidFill><a:ln><a:solidFill><a:srgbClr val="FFFFFF"/></a:solidFill>'
                         f'</a:ln></c:spPr></c:marker>'))
    fmt = NUMFMT["int"] if ln_top >= 100 or all(float(v).is_integer() for v in ln_vals) else NUMFMT["dec1"]
    ser.append(parse_xml(f'<c:dLbls {nsdecls("c", "a")}><c:numFmt formatCode="{fmt.replace(chr(34), "&quot;")}" sourceLinked="0"/>'
                         f'<c:spPr><a:noFill/><a:ln><a:noFill/></a:ln></c:spPr><c:txPr><a:bodyPr/><a:lstStyle/><a:p><a:pPr>'
                         f'<a:defRPr sz="900" b="1"><a:solidFill><a:srgbClr val="0F1F3D"/></a:solidFill>'
                         f'<a:latin typeface="{FONT}"/></a:defRPr></a:pPr><a:endParaRPr lang="en-US"/></a:p></c:txPr>'
                         f'<c:dLblPos val="t"/><c:showLegendKey val="0"/><c:showVal val="1"/><c:showCatName val="0"/>'
                         f'<c:showSerName val="0"/><c:showPercent val="0"/><c:showBubbleSize val="0"/></c:dLbls>'))
    ser.append(copy.deepcopy(src.find(qn("c:cat"))))
    ser.append(copy.deepcopy(src.find(qn("c:val"))))
    ser.append(parse_xml(f'<c:smooth {nsdecls("c")} val="0"/>'))
    ax_ids = [a.get("val") for a in bar.findall(qn("c:axId"))]
    if c.secondary:
        ax_ids = ["60001", "60002"]
    lc = parse_xml(f'<c:lineChart {nsdecls("c")}><c:grouping val="standard"/><c:varyColors val="0"/></c:lineChart>')
    lc.append(ser)
    lc.append(parse_xml(f'<c:marker {nsdecls("c")} val="1"/>'))
    for a in ax_ids:
        lc.append(parse_xml(f'<c:axId {nsdecls("c")} val="{a}"/>'))
    bar.addnext(lc)
    if c.secondary:
        base = max(0.0, ln_lo - (ln_top - ln_lo) * 1.8) if ln_top > ln_lo else 0.0
        mx = ln_top + (ln_top - base) * 0.2 if ln_top > base else ln_top * 1.2 + 1
        last_ax = pa.findall(qn("c:valAx"))[-1]
        cat2 = parse_xml(
            f'<c:catAx {nsdecls("c")}><c:axId val="60001"/><c:scaling><c:orientation val="minMax"/></c:scaling>'
            f'<c:delete val="1"/><c:axPos val="b"/><c:majorTickMark val="none"/><c:minorTickMark val="none"/>'
            f'<c:tickLblPos val="nextTo"/><c:crossAx val="60002"/><c:crosses val="autoZero"/><c:auto val="1"/>'
            f'<c:lblAlgn val="ctr"/><c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>')
        val2 = parse_xml(
            f'<c:valAx {nsdecls("c", "a")}><c:axId val="60002"/><c:scaling><c:orientation val="minMax"/>'
            f'<c:max val="{mx:.4f}"/><c:min val="{base:.4f}"/></c:scaling><c:delete val="0"/><c:axPos val="r"/>'
            f'<c:numFmt formatCode="#,##0" sourceLinked="0"/><c:majorTickMark val="none"/><c:minorTickMark val="none"/>'
            f'<c:tickLblPos val="nextTo"/><c:spPr><a:ln><a:noFill/></a:ln></c:spPr><c:txPr><a:bodyPr/><a:lstStyle/>'
            f'<a:p><a:pPr><a:defRPr sz="900"><a:solidFill><a:srgbClr val="5E6B85"/></a:solidFill>'
            f'<a:latin typeface="{FONT}"/></a:defRPr></a:pPr><a:endParaRPr lang="en-US"/></a:p></c:txPr>'
            f'<c:crossAx val="60001"/><c:crosses val="max"/><c:crossBetween val="between"/></c:valAx>')
        last_ax.addnext(cat2)
        cat2.addnext(val2)
    _legend(ch)
    return ch


def _bar(slide, c: Chart, x, y, w, h, horizontal: bool):
    kind = XL_CHART_TYPE.BAR_CLUSTERED if horizontal else XL_CHART_TYPE.COLUMN_CLUSTERED
    cats = [str(v) for v in c.cats]
    if horizontal:
        cats = [s if len(s) <= 26 else s[:25] + "…" for s in map(str, cats)]
    ch = _add_chart(slide, kind, x, y, w, h, cats, c.series)
    plot = ch.plots[0]
    nc = len(cats)
    plot.gap_width = (180 if nc <= 3 else 110 if nc <= 5 else 60) if horizontal else (120 if nc <= 3 else 70)
    s = plot.series[0]
    s.format.fill.solid()
    s.format.fill.fore_color.rgb = rgb(NAVY if horizontal else TEAL)
    s.invert_if_negative = False
    _labels(plot, c.fmt)
    top = _top(c.series[0][1])
    ch.value_axis.minimum_scale = 0
    ch.value_axis.maximum_scale = top * (1.22 if horizontal else 1.18)
    if horizontal:
        _axis_clean(ch.value_axis, show_labels=False, grid=False)
        _axis_clean(ch.category_axis, grid=False)
        ch.category_axis.reverse_order = True
        ch.category_axis.tick_labels.font.color.rgb = rgb(INK)
    else:
        _axis_clean(ch.value_axis)
        _axis_clean(ch.category_axis, grid=False)
    ch.has_legend = False
    return ch


def _line(slide, c: Chart, x, y, w, h):
    ch = _add_chart(slide, XL_CHART_TYPE.LINE_MARKERS, x, y, w, h, c.cats, c.series)
    plot = ch.plots[0]
    s = plot.series[0]
    s.smooth = False
    s.format.line.color.rgb = rgb(TEAL)
    s.format.line.width = Pt(2.25)
    s.marker.style = 8  # circle
    s.marker.size = 7
    s.marker.format.fill.solid()
    s.marker.format.fill.fore_color.rgb = rgb(TEAL)
    _labels(plot, c.fmt, XL_LABEL_POSITION.ABOVE)
    _axis_clean(ch.value_axis)
    _axis_clean(ch.category_axis, grid=False)
    ch.value_axis.minimum_scale = 0
    ch.value_axis.maximum_scale = _top(s.values) * 1.25
    ch.has_legend = False
    return ch


def _manual_layout(x, y, w, h, target=True):
    t = '<c:layoutTarget val="inner"/>' if target else ""
    return (f'<c:layout {nsdecls("c")}><c:manualLayout>{t}<c:xMode val="edge"/><c:yMode val="edge"/>'
            f'<c:x val="{x}"/><c:y val="{y}"/><c:w val="{w}"/><c:h val="{h}"/></c:manualLayout></c:layout>')


def _donut(slide, c: Chart, x, y, w, h):
    vals = c.series[0][1]
    tot = sum(v or 0 for v in vals) or 1
    share = [(v or 0) / tot * 100 for v in vals]
    cats = list(c.cats)
    ch = _add_chart(slide, XL_CHART_TYPE.DOUGHNUT, x, y, w, h, cats, [(c.series[0][0], vals)])
    plot = ch.plots[0]
    hole = plot._element.find(qn("c:holeSize"))
    if hole is not None:
        hole.set("val", "60")
    for i, pt in enumerate(plot.series[0].points):
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(PALETTE[i % len(PALETTE)])
        pt.format.line.color.rgb = rgb("#FFFFFF")
        pt.format.line.width = Pt(1.5)
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.show_percentage, dl.show_value, dl.show_category_name = True, False, False
    dl.number_format = "0%"
    dl.number_format_is_linked = False
    _font(dl, 9, "#FFFFFF", True)
    for i, pt in enumerate(plot.series[0].points):
        if share[i] < 6:
            pt.data_label.has_text_frame = True
            pt.data_label.text_frame.text = ""
    ch.has_legend = False
    pa = ch._chartSpace.chart.plotArea
    old = pa.find(qn("c:layout"))
    if old is not None:
        pa.remove(old)
    pa.insert(0, parse_xml(_manual_layout(0.02, 0.04, 0.5, 0.92)))
    # legend drawn as text so it looks identical in PowerPoint, Keynote and LibreOffice
    lx, lw = x + 0.56 * w, 0.43 * w
    rowh = min(0.42, (0.84 * h) / max(len(cats), 1))
    ly = y + (h - rowh * len(cats)) / 2
    for i, (nme, v, sh) in enumerate(zip(c.cats, vals, share)):
        yy = ly + i * rowh
        _box(slide, lx, yy + rowh / 2 - 0.07, 0.14, 0.14, fill=PALETTE[i % len(PALETTE)], line=None, radius=0.2)
        nm = nme if len(nme) <= 20 else nme[:19] + "…"
        _text(slide, lx + 0.24, yy, lw - 0.24, rowh,
              [[(nm, {"color": INK, "size": 10.5}), (f"   {sh:.0f}%", {"bold": True, "color": INK, "size": 10.5}),
                (f" ({(v or 0):,.0f})", {"color": MUTED, "size": 10})]], anchor=MSO_ANCHOR.MIDDLE)
    # total in the centre
    d = min(0.5 * w, 0.92 * h)
    cx, cy = x + (0.02 + 0.25) * w, y + (0.04 + 0.46) * h
    _text(slide, cx - d * 0.32, cy - 0.32, d * 0.64, 0.64,
          [[(f"{tot:,.0f}", {"size": 18, "bold": True, "color": INK})]] + ([[(c.center, {"size": 9, "color": MUTED})]] if c.center else []),
          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return ch


def _funnel(slide, c: Chart, x, y, w, h):
    stages, vals = c.cats, [v or 0 for v in c.series[0][1]]
    n = len(stages)
    if not n:
        return
    first = vals[0] or 1
    mx = max(vals) or 1
    rh = min(0.62, h / n)
    gap = rh * 0.16
    name_w, val_w = w * 0.26, w * 0.2
    bar_x0, bar_w = x + name_w, w - name_w - val_w
    top = y + (h - rh * n) / 2
    for i, (s, v) in enumerate(zip(stages, vals)):
        bw = max(0.35, bar_w * v / mx)
        by = top + i * rh
        col = blend("#3CC1C0", NAVY, i / max(n - 1, 1))
        _box(slide, bar_x0 + (bar_w - bw) / 2, by + gap / 2, bw, rh - gap, fill=col, line=None, radius=0.15)
        _text(slide, x, by, name_w - 0.1, rh, str(s), 10.5, INK, True, PP_ALIGN.RIGHT, MSO_ANCHOR.MIDDLE)
        _text(slide, x + w - val_w + 0.08, by, val_w - 0.08, rh,
              [[(f"{v:,.0f}", {"bold": True, "color": INK, "size": 11}), (f"  {v / first * 100:.0f}%", {"color": MUTED, "size": 10})]],
              anchor=MSO_ANCHOR.MIDDLE)


def draw_chart(slide, c: Chart, x, y, w, h):
    """Card with title + chart (or a 'not available' note)."""
    _card(slide, x, y, w, h, c.title, c.subtitle)
    cx, cy, cw, chh = x + 0.15, y + 0.82, w - 0.3, h - 1.0 - (0.25 if c.note else 0)
    if not c.ok:
        b = _box(slide, x + 0.3, y + 1.0, w - 0.6, 1.0, fill="#F7F9FC", line="#C8D1E0")
        b.line.dash_style = 4
        _text(slide, x + 0.45, y + 1.05, w - 0.9, 0.9, f"Not available — {c.reason}", 11, MUTED, anchor=MSO_ANCHOR.MIDDLE)
        return
    if c.kind == "combo":
        _combo(slide, c, cx, cy, cw, chh)
    elif c.kind in ("bar", "hbar"):
        _bar(slide, c, cx, cy, cw, chh, c.kind == "hbar")
    elif c.kind == "line":
        _line(slide, c, cx, cy, cw, chh)
    elif c.kind == "donut":
        _donut(slide, c, cx, cy, cw, chh)
    elif c.kind == "funnel":
        _funnel(slide, c, cx + 0.1, cy + 0.05, cw - 0.2, chh - 0.1)
    if c.note:
        _text(slide, x + 0.22, y + h - 0.34, w - 0.44, 0.24, c.note, 9, MUTED, italic=True)


# ------------------------------------------------------------------ slides
def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _stripes(slide, x, y, h, color=GREEN, n=2):
    for i in range(n):
        s = slide.shapes.add_shape(MSO_SHAPE.PARALLELOGRAM, Inches(x + i * 0.62), Inches(y), Inches(0.9), Inches(h))
        s.adjustments[0] = 0.45
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(color)
        s.line.fill.background()
        s.shadow.inherit = False


def cover(prs, title: str, area: str, rep: Report, lines: list[str] | None = None):
    s = _blank(prs)
    _bg(s, "#FFFFFF")
    panel = _box(s, 8.75, 0, W - 8.75, H, fill=NAVY, line=None, shape=MSO_SHAPE.RECTANGLE)
    _stripes(s, 10.0, 1.6, 4.3)
    _text(s, 9.15, H - 1.25, 3.9, 0.4, "HR Analytics", 16, "#FFFFFF", True)
    _text(s, 9.15, H - 0.85, 3.9, 0.35, "Automated from your HR workbook", 11, "#C9D6EE")
    _logo(s, 0.8, 0.85, 0.85)
    _box(s, 0.8, 2.15, 1.2, 0.06, fill=GREEN, line=None, shape=MSO_SHAPE.RECTANGLE)
    _text(s, 0.8, 2.4, 7.6, 0.9, title, 38, NAVY, True)
    _text(s, 0.8, 3.3, 7.6, 0.45, area, 18, MUTED)
    _text(s, 0.8, 3.8, 7.6, 0.5, TAGLINE, 20, GREEN, True)
    loc = rep.location if rep.location and rep.location != "All" else "All locations"
    meta = lines or [f"Reporting period:  {rep.period}  ({rep.period_range})", f"Location:  {loc}",
                     f"Source workbook:  {rep.source or '—'}", f"Generated:  {rep.generated:%d %b %Y, %H:%M}"]
    _text(s, 0.8, 4.75, 7.6, 1.6, [[(m, {})] for m in meta], 12, INK)
    _text(s, 0.8, H - 0.6, 7.6, 0.3, f"{APP_NAME}  ·  Confidential", 9, MUTED)
    return s


def divider(prs, rep: Report, n: int):
    s = _blank(prs)
    _bg(s, NAVY)
    _stripes(s, 10.4, 0, H, "#1F5BB8")
    _stripes(s, 11.1, 2.2, 3.1, GREEN)
    _text(s, 0.9, 2.6, 9, 0.9, rep.title, 36, "#FFFFFF", True)
    _text(s, 0.9, 3.5, 9, 0.5, rep.area, 18, "#C9D6EE")
    _text(s, 0.9, 4.1, 9, 0.4, f"{rep.period}  ·  {rep.period_range}", 12, "#C9D6EE")
    _text(s, W - 1.45, H - 0.38, 1.0, 0.25, str(n), 9, "#C9D6EE", align=PP_ALIGN.RIGHT)


def kpi_slide(prs, rep: Report, n: int):
    s = _blank(prs)
    _header(s, rep)
    _text(s, 0.45, 1.3, 8, 0.35, "Key Metrics", 15, INK, True)
    vs = next((d.vs for _, d in rep.kpis if d), "")
    if vs:
        _text(s, W - 5.45, 1.32, 5.0, 0.3, f"Arrows compare with {vs}. Green = better, red = worse, blue = neutral.",
              9.5, MUTED, align=PP_ALIGN.RIGHT)
    head, extra = rep.kpis[:6], rep.kpis[6:10]
    cw, ch_, gx, gy = (W - 0.9 - 0.5) / 3, 1.62, 0.25, 0.22
    for i, (m, d) in enumerate(head):
        col, row = i % 3, i // 3
        x, y = 0.45 + col * (cw + gx), 1.78 + row * (ch_ + gy)
        _box(s, x, y, cw, ch_)
        _box(s, x, y + 0.18, 0.07, ch_ - 0.36, fill=TEAL if i % 2 else NAVY, line=None, shape=MSO_SHAPE.RECTANGLE)
        _text(s, x + 0.3, y + 0.16, cw - 0.5, 0.3, P.label(m), 12, MUTED, True)
        if m.available:
            _text(s, x + 0.3, y + 0.45, cw - 0.5, 0.6, P.fmt(m), 30, INK, True)
            if d:
                tone = {"good": GOOD, "bad": BAD}.get(d.tone, NAVY)
                arrow = "" if d.value == 0 else ("▲ " if d.up else "▼ ")
                _text(s, x + 0.3, y + 1.12, cw - 0.5, 0.3,
                      [[(arrow + d.text, {"bold": True, "color": tone, "size": 12}), (f"   vs {d.vs}", {"color": MUTED, "size": 10})]])
            elif m.detail:
                _text(s, x + 0.3, y + 1.14, cw - 0.5, 0.3, m.detail, 9.5, MUTED)
        else:
            _text(s, x + 0.3, y + 0.5, cw - 0.5, 0.4, "Not available", 18, "#7A8499", True)
            _text(s, x + 0.3, y + 0.95, cw - 0.5, 0.55, m.reason.replace("Not available — ", ""), 9.5, MUTED)
    if extra:
        y = 1.78 + 2 * (ch_ + gy) + 0.05
        mw = (W - 0.9 - 0.25 * (len(extra) - 1)) / len(extra) if len(extra) > 1 else 4.0
        for i, (m, d) in enumerate(extra):
            x = 0.45 + i * (mw + 0.25)
            _box(s, x, y, mw, 1.05)
            _text(s, x + 0.22, y + 0.12, mw - 0.4, 0.26, P.label(m), 10.5, MUTED, True)
            val = P.fmt(m) if m.available else "Not available"
            runs = [(val, {"size": 18 if m.available else 13, "bold": True, "color": INK if m.available else "#7A8499"})]
            if d and m.available:
                tone = {"good": GOOD, "bad": BAD}.get(d.tone, NAVY)
                runs.append(("   " + ("" if d.value == 0 else ("▲ " if d.up else "▼ ")) + d.text, {"size": 10.5, "bold": True, "color": tone}))
            _text(s, x + 0.22, y + 0.42, mw - 0.4, 0.5, [runs], anchor=MSO_ANCHOR.MIDDLE)
    _footer(s, rep, n)


def chart_slide(prs, rep: Report, charts: list[Chart], n: int):
    s = _blank(prs)
    _header(s, rep)
    if len(charts) == 1:
        draw_chart(s, charts[0], 0.45, 1.35, W - 0.9, 5.55)
    else:
        cw = (W - 0.9 - 0.25) / 2
        for i, c in enumerate(charts[:2]):
            draw_chart(s, c, 0.45 + i * (cw + 0.25), 1.35, cw, 5.55)
    _footer(s, rep, n)


def table_slide(prs, rep: Report, n: int):
    title, df = rep.table
    s = _blank(prs)
    _header(s, rep)
    _card(s, 0.45, 1.35, W - 0.9, 5.55, title.split(" — ")[0], title.split(" — ")[1] if " — " in title else "")
    rows, cols = len(df) + 1, len(df.columns)
    gt = s.shapes.add_table(rows, cols, Inches(0.7), Inches(2.2), Inches(W - 1.4), Inches(min(4.4, 0.34 * rows))).table
    for j, c in enumerate(df.columns):
        cell = gt.cell(0, j)
        cell.text = str(c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(NAVY)
        p = cell.text_frame.paragraphs[0]
        _font(p, 10.5, "#FFFFFF", True)
    for i, (_, r) in enumerate(df.iterrows(), start=1):
        for j, c in enumerate(df.columns):
            v = r[c]
            txt = "" if v is None or (isinstance(v, float) and v != v) else (P.num(float(v)) if isinstance(v, float) else str(v))
            cell = gt.cell(i, j)
            cell.text = txt
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb("#FFFFFF" if i % 2 else "#F3F6FA")
            _font(cell.text_frame.paragraphs[0], 10, INK)
    _footer(s, rep, n)


def _focus_rows(rep: Report):
    rows = []
    for i in rep.focus:
        rows.append(("⚠" if i.kind == "attention" else "✔", "#D08A14" if i.kind == "attention" else GREEN, i.text))
        if i.action:
            rows.append(("➜", TEAL, i.action))
    return rows or [("•", MUTED, "No rule-based items from the current data.")]


def _focus_card(s, rep: Report, x, y, w, h, max_rows=9):
    _card(s, x, y, w, h, "Initiatives & Focus", "Rule-based items from the current data")
    yy = y + 0.85
    for sym, col, txt in _focus_rows(rep)[:max_rows]:
        lines = max(1, int(len(txt) * 0.085 / max(w - 1.0, 1)) + 1)
        _text(s, x + 0.3, yy, 0.35, 0.4, sym, 14, col, True)
        _text(s, x + 0.7, yy + 0.03, w - 1.0, 0.3 * lines, txt, 11.5, INK)
        yy += 0.22 + 0.27 * lines
        if yy > y + h - 0.5:
            break


def _notes_text(rep: Report):
    notes = [[(f"{t}: ", {"bold": True, "color": INK}), (r, {"color": MUTED})] for t, r in rep.notes[:7]]
    notes.append([("Confidentiality: ", {"bold": True, "color": INK}), (MIN_GROUP_NOTE, {"color": MUTED})])
    vs = next((d.vs for _, d in rep.kpis if d), "")
    if vs:
        notes.append([("Comparison: ", {"bold": True, "color": INK}), (f"changes are vs {vs}.", {"color": MUTED})])
    return notes


def focus_slide(prs, rep: Report, n: int, chart: Chart | None = None):
    """Closing slide: Initiatives & Focus with Data Notes — or, when one chart is left over, chart + focus."""
    s = _blank(prs)
    _header(s, rep)
    if chart is not None:
        cw = (W - 0.9 - 0.25) / 2
        draw_chart(s, chart, 0.45, 1.35, cw, 5.55)
        rx = 0.45 + cw + 0.25
        _focus_card(s, rep, rx, 1.35, cw, 3.6, max_rows=6)
        _card(s, rx, 5.15, cw, 1.75, "Data Notes")
        _text(s, rx + 0.22, 5.62, cw - 0.44, 1.2, _notes_text(rep)[-3:], 9.5)
    else:
        lw = (W - 0.9) * 0.6
        _focus_card(s, rep, 0.45, 1.35, lw, 5.55)
        rx, rw = 0.45 + lw + 0.25, W - 0.9 - lw - 0.25
        _card(s, rx, 1.35, rw, 5.55, "Data Notes", "What could not be shown, and why")
        _text(s, rx + 0.25, 2.15, rw - 0.5, 4.5, _notes_text(rep), 10.5)
    _footer(s, rep, n)


def _report_slides(prs, rep: Report, n: int, with_cover: bool) -> int:
    if with_cover:
        cover(prs, rep.title, rep.area, rep)
    else:
        divider(prs, rep, n)
    n += 1
    if not rep.available:
        s = _blank(prs)
        _header(s, rep)
        _text(s, 0.8, 2.5, W - 1.6, 1, f"Not available — {rep.reason}", 16, MUTED)
        _footer(s, rep, n)
        return n + 1
    kpi_slide(prs, rep, n)
    n += 1
    ok = [c for c in rep.charts if c.ok]
    last = ok.pop() if len(ok) % 2 else None
    for i in range(0, len(ok), 2):
        chart_slide(prs, rep, ok[i:i + 2], n)
        n += 1
    if rep.table is not None:
        table_slide(prs, rep, n)
        n += 1
    focus_slide(prs, rep, n, last)
    return n + 1


def _new() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    return prs


def render(rep: Report) -> bytes:
    prs = _new()
    _report_slides(prs, rep, 1, True)
    prs.core_properties.title = f"SWITCH {rep.title} — {rep.period}"
    prs.core_properties.author = "SWITCH HR Analytics"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def render_pack(reports: list[Report]) -> bytes:
    """One deck with every dashboard: cover, contents, then a section per dashboard."""
    prs = _new()
    first = reports[0]
    cover(prs, "HR Analytics Pack", "All dashboards", first)
    s = _blank(prs)
    _header(s, first, "Contents")
    _text(s, 0.45, 1.35, 8, 0.4, "Contents", 18, INK, True)
    for i, r in enumerate(reports):
        y = 2.0 + i * 0.62
        _box(s, 0.45, y, 0.46, 0.46, fill=NAVY if i % 2 == 0 else TEAL, line=None, radius=0.3)
        _text(s, 0.45, y, 0.46, 0.46, str(i + 1), 13, "#FFFFFF", True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        _text(s, 1.1, y + 0.02, 6, 0.42, [[(r.title, {"bold": True, "color": INK, "size": 14}),
                                           (f"   {r.area}", {"color": MUTED, "size": 11})]], anchor=MSO_ANCHOR.MIDDLE)
    _footer(s, first, 2)
    n = 3
    for r in reports:
        n = _report_slides(prs, r, n, False)
    prs.core_properties.title = f"SWITCH HR Analytics Pack — {first.period}"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
