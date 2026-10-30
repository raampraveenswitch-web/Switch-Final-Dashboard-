"""Accuracy check. Builds a messy FICTIONAL workbook, computes the expected answers with an independent
plain-Python oracle (no app code), loads the workbook through the app and compares every metric.
Usage: python tools/accuracy_check.py [out_dir]"""
import random
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from services import metrics as M  # noqa: E402
from services.app_state import Store  # noqa: E402

rnd = random.Random(2026)
END = date(2026, 9, 30)
START = date(2025, 10, 1)
P = M.Period(pd.Timestamp(START), pd.Timestamp(END))
dt = lambda x: datetime(x.year, x.month, x.day)  # noqa: E731
DEPTS = ["Engineering", "Operations", "Sales", "Finance"]

# ---------------------------------------------------------------- truth records
emps = []
for i in range(1, 241):
    j = END - timedelta(days=rnd.randint(20, 2400))
    x = None
    if rnd.random() < 0.2:
        c = j + timedelta(days=rnd.randint(30, 1200))
        x = c if c <= END + timedelta(days=60) else None
    emps.append(dict(id=f"EMP{i:04d}", dept=rnd.choice(DEPTS), doj=j, lwd=x,
                     typ=rnd.choice(["Resignation", "Termination"]) if x else None))
reqs = []
for i in range(1, 121):
    o = END - timedelta(days=rnd.randint(1, 700))
    r = rnd.random()
    st, f = ("Open - Sourcing" if r < .25 else "Cancelled" if r < .32 else "On Hold" if r < .38 else "Closed - Joined"), None
    if st == "Closed - Joined":
        f = o + timedelta(days=rnd.randint(15, 100))
        if f > END:
            st, f = "Open", None
    reqs.append(dict(id=f"R{i:04d}", o=o, f=f, st=st,
                     offer=(rnd.choice(["Accepted", "Accepted", "Declined", "Rejected"]) if r > .38 else None),
                     cost=rnd.randint(10, 80) * 1000 if st == "Closed - Joined" else None))
reqs[5].update(st="Closed - Joined", o=END - timedelta(days=50), f=END - timedelta(days=70), cost=20000)  # filled BEFORE open: must be excluded from time-to-hire
trn = []
for _ in range(500):
    h = rnd.choice([2, 4, 8])
    trn.append(dict(emp=rnd.choice(emps)["id"], hrs=h, d=END - timedelta(days=rnd.randint(1, 700)),
                    st=rnd.choice(["Completed"] * 6 + ["Not Completed"]), cert=rnd.choice(["Yes", "No", None]),
                    pre=rnd.randint(40, 70), post=rnd.randint(55, 95), cost=rnd.randint(2, 30) * 100))
trn[3]["hrs"] = -4          # invalid: must be excluded, not counted
pms = []
for e in emps[:180]:
    ok = rnd.random() < .85
    g = rnd.randint(3, 6)
    pms.append(dict(id=e["id"], rating=rnd.choice([2, 3, 3, 4, 4, 5]) if ok else None, st="Completed" if ok else "Pending",
                    gs=g if ok else None, ga=rnd.randint(0, g) if ok else None, idp=rnd.choice(["Yes", "No"])))
cases = []
for i in range(1, 81):
    o = END - timedelta(days=rnd.randint(1, 600))
    cl = rnd.random() < .7
    cases.append(dict(id=f"CS{i:03d}", o=o, st="Resolved" if cl else rnd.choice(["Open", "Pending"]),
                      c=o + timedelta(days=rnd.randint(2, 45)) if cl else None))

# -------------------------------------------------------------------- workbook (messy on purpose)
wb = Workbook(); wb.remove(wb.active)
north, south = emps[:120], emps[120:]
ws = wb.create_sheet("Emp - North"); ws.append(["SWITCH HR — Employee Master (North)"]); ws.append([])
ws.append(["Emp Code", "Dept", "DOJ", "LWD", "Separation Type"])
for e in north:
    ws.append([e["id"], e["dept"], dt(e["doj"]), dt(e["lwd"]) if e["lwd"] else None, e["typ"]])
for e in north[:3]:                                  # 3 exact duplicate rows
    ws.append([e["id"], e["dept"], dt(e["doj"]), dt(e["lwd"]) if e["lwd"] else None, e["typ"]])
ws = wb.create_sheet("Emp - South"); ws.append(["Employee ID", "Department", "Date of Joining", "Date of Exit", "Exit Type"])
for k, e in enumerate(south):
    doj = e["doj"].strftime("%d/%m/%Y") if k % 7 == 0 else dt(e["doj"])     # some dates stored as dd/mm/yyyy text
    ws.append([e["id"], e["dept"], doj, dt(e["lwd"]) if e["lwd"] else None, e["typ"]])
ws.append(["EMPX1", "Sales", dt(END - timedelta(days=100)), "TBD", None])      # invalid exit date
ws = wb.create_sheet("Hiring")
ws.append(["Req No", "Position", "Req Status", "Requisition Open Date", "Offer Accepted Date", "Offer Outcome", "Recruitment Cost"])
for r in reqs:
    ws.append([r["id"], "Role", r["st"], dt(r["o"]), dt(r["f"]) if r["f"] else None, r["offer"],
               f"₹{r['cost']:,}" if r["cost"] else None])               # currency stored as text
ws = wb.create_sheet("L&D"); ws.append(["Training log"]); 
ws.append(["Emp ID", "Course Name", "Training Hours", "Date", "Training Status", "Certified", "Pre Score", "Post Score", "Training Cost"])
for t in trn:
    ws.append([t["emp"], "Course", t["hrs"], dt(t["d"]), t["st"], t["cert"], t["pre"], t["post"], t["cost"]])
ws = wb.create_sheet("PMS"); ws.append(["Employee Code", "Final Rating", "Appraisal Status", "Goals Set", "Goals Achieved", "IDP Created"])
for p_ in pms:
    ws.append([p_["id"], p_["rating"], p_["st"], p_["gs"], p_["ga"], p_["idp"]])
ws = wb.create_sheet("Grievances"); ws.append(["Case Number", "Case Status", "Date Raised", "Date Closed"])
for c in cases:
    ws.append([c["id"], c["st"], dt(c["o"]), dt(c["c"]) if c["c"] else None])
wb.create_sheet("Scratch")
out = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
path = out / "SWITCH_accuracy_test_workbook_FICTIONAL.xlsx"; wb.save(path)

# ------------------------------------------------------------------ independent oracle (plain python)
# EMPX1 was written above with an unreadable exit date ("TBD"); the documented rule treats an invalid date as
# "no exit date", so the person counts as active. The oracle must include that row.
emps.append(dict(id="EMPX1", dept="Sales", doj=END - timedelta(days=100), lwd=None, typ=None))
inp = lambda d_: START <= d_ <= END  # noqa: E731
def active(d_): return {e["id"] for e in emps if e["doj"] <= d_ and (e["lwd"] is None or e["lwd"] > d_)}
head_end, head_open = len(active(END)), len(active(START - timedelta(days=1)))
exits = [e for e in emps if e["lwd"] and inp(e["lwd"])]
filled = [r for r in reqs if r["st"] == "Closed - Joined" and r["f"] and inp(r["f"])]
ttd = [(r["f"] - r["o"]).days for r in filled if (r["f"] - r["o"]).days >= 0]
offers = [r["offer"] for r in reqs if inp(r["o"]) and r["offer"]]
costed = [r["cost"] for r in filled if r["cost"]]
openr = [r for r in reqs if r["st"].startswith("Open") and r["o"] <= END]
tin = [t for t in trn if inp(t["d"])]
tv = lambda t: t["hrs"] if t["hrs"] >= 0 else 0  # noqa: E731
done = {t["emp"] for t in tin if t["st"] == "Completed"} & active(END)
certs = {}
for t in tin:
    if t["cert"]:
        certs[t["emp"]] = certs.get(t["emp"], False) or t["cert"] == "Yes"
rated = [p_["rating"] for p_ in pms if p_["rating"]]
gl = [(p_["gs"], p_["ga"]) for p_ in pms if p_["gs"]]
cin = [c for c in cases if inp(c["o"])]
cclosed = [c for c in cin if c["st"] == "Resolved"]
exp = {
    "emp_headcount": head_end, "emp_new_hires": sum(inp(e["doj"]) for e in emps), "emp_exits": len(exits),
    "emp_attrition": len(exits) / ((head_open + head_end) / 2) * 100,
    "emp_tenure": sum((END - e["doj"]).days / 365.25 for e in emps if e["id"] in active(END)) / head_end,
    "rec_requisitions": sum(inp(r["o"]) for r in reqs), "rec_filled": len(filled), "rec_new_hires": len(filled),
    "rec_time_to_hire": sum(ttd) / len(ttd),
    "rec_offer_acceptance": offers.count("Accepted") / len(offers) * 100, "rec_cost_per_hire": sum(costed) / len(costed),
    "rec_open_count": len(openr), "rec_open_ageing": sum((END - r["o"]).days for r in openr) / len(openr),
    "trn_hours": sum(tv(t) for t in tin), "trn_employees": len({t["emp"] for t in tin if t["st"] == "Completed"}),
    "trn_participation": len(done) / head_end * 100, "trn_avg_hours": sum(tv(t) for t in tin) / head_end,
    "trn_cert_rate": sum(certs.values()) / len(certs) * 100,
    "trn_effectiveness": sum(t["post"] - t["pre"] for t in tin) / len(tin), "trn_cost": sum(t["cost"] for t in tin),
    "pms_eligible": len(pms), "pms_completed": sum(p_["st"] == "Completed" for p_ in pms),
    "pms_completion": sum(p_["st"] == "Completed" for p_ in pms) / len(pms) * 100, "pms_avg_rating": sum(rated) / len(rated),
    "pms_goal_rate": sum(a for _, a in gl) / sum(g for g, _ in gl) * 100, "pms_idp": sum(p_["idp"] == "Yes" for p_ in pms),
    "rel_cases": len(cin), "rel_closed": len(cclosed), "rel_resolution_rate": len(cclosed) / len(cin) * 100,
    "rel_resolution_days": sum((c["c"] - c["o"]).days for c in cclosed) / len(cclosed),
}

# --------------------------------------------------------------------------- app under test
s = Store(); ok = s.ingest_bytes(path.read_bytes(), path.name)
res = {k: v for k, v in __import__("services.overview", fromlist=["x"]).all_metrics(s.model.datasets, s.settings, P).items()}
rows, bad = [], 0
for k, e in exp.items():
    got = res[k].value if k in res else None
    good = got is not None and abs(got - e) < 1e-6 * max(1, abs(e))
    bad += not good
    rows.append((k, round(e, 4), None if got is None else round(got, 4), "PASS" if good else "FAIL"))
cl = {n: (c.dataset, c.confidence) for n, c in s.model.classifications.items()}
print("State:", s.state.value, "| loaded:", ok)
print("Sheet classification:", cl)
print("Warnings:", [i.message for i in s.model.issues])
print(f"\n{'metric':24}{'expected':>14}{'app':>14}  result")
for r in rows:
    print(f"{r[0]:24}{r[1]:>14}{str(r[2]):>14}  {r[3]}")
print(f"\n{len(rows) - bad}/{len(rows)} metrics match the independent calculation")

# ------------------------------------------------ edit test: remove one active employee -> headcount must drop by 1
wb2 = __import__("openpyxl").load_workbook(path)
ws2 = wb2["Emp - South"]
victim = next(r for r in ws2.iter_rows(min_row=2) if r[3].value is None and r[0].value != "EMPX1")
ws2.delete_rows(victim[0].row)
p2 = out / "edited.xlsx"; wb2.save(p2)
assert s.ingest_bytes(p2.read_bytes(), "edited.xlsx")
h2 = M.emp_headcount(s.model.datasets["employee"], pd.Timestamp(END)).value
print(f"Replace/recalculate test: headcount {head_end} -> {int(h2)} (expected {head_end - 1}):", "PASS" if h2 == head_end - 1 else "FAIL")
p2.unlink()
sys.exit(1 if bad else 0)
