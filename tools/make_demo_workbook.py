"""Generates a FICTIONAL workbook for trying the dashboard. Standalone: the app never imports this.
Usage: python tools/make_demo_workbook.py [output.xlsx]"""
import random
import sys
from datetime import date, timedelta

from openpyxl import Workbook

rnd = random.Random(7)
END = date(2026, 9, 30)
DEPTS = ["Engineering", "Operations", "Sales", "Finance", "HR"]
LOCS = ["Chennai", "Pune", "Bengaluru"]
TYPES = ["Permanent", "Permanent", "Permanent", "Contract", "Intern"]


def d(days_back):
    return END - timedelta(days=days_back)


wb = Workbook()
ws = wb.active
ws.title = "About"
ws.append(["FICTIONAL DEMO DATA — generated for testing only. Not real employees, requisitions or results."])

# Employees
emp = wb.create_sheet("Employees")
emp.append(["Employee ID", "Department", "Location", "Employee Type", "Gender", "Grade", "Date of Joining",
            "Date of Exit", "Exit Type", "Exit Reason"])
for i in range(1, 161):
    join = d(rnd.randint(30, 1800))
    left, et, er = None, None, None
    if rnd.random() < 0.18:
        left = join + timedelta(days=rnd.randint(60, 900))
        if left <= END:
            et = rnd.choice(["Resignation", "Resignation", "Resignation", "Termination"])
            er = rnd.choice(["Better opportunity", "Relocation", "Higher studies", "Performance", "Personal"])
        else:
            left = None
    emp.append([f"DEMO{i:03d}", rnd.choice(DEPTS), rnd.choice(LOCS), rnd.choice(TYPES), rnd.choice(["F", "M", "M", "F"]),
                rnd.choice(["G1", "G2", "G3", "G4"]), join, left, et, er])

# Recruitment (requisition level)
rec = wb.create_sheet("Hiring")
rec.append(["Req ID", "Job Title", "Department", "Location", "Priority", "Recruiter", "Hiring Manager", "Status",
            "Requisition Open Date", "Filled Date", "Source", "Offer Status", "Cost per Hire", "Stage"])
titles = ["Engineer", "Analyst", "Technician", "Manager", "Executive", "Designer"]
stages = ["Applied", "Screened", "Interviewed", "Offered", "Joined"]
for i in range(1, 91):
    opened = d(rnd.randint(5, 540))
    r = rnd.random()
    if r < 0.62:
        status, filled = "Filled", opened + timedelta(days=rnd.randint(20, 90))
        if filled > END:
            status, filled = "Open", None
    elif r < 0.9:
        status, filled = "Open", None
    else:
        status, filled = "Cancelled", None
    offer = rnd.choice(["Accepted", "Accepted", "Declined"]) if status == "Filled" else None
    if status == "Filled" and offer == "Declined":
        offer = "Accepted"
    stage = "Joined" if status == "Filled" else rnd.choice(stages[:4])
    rec.append([f"REQ{i:03d}", rnd.choice(titles), rnd.choice(DEPTS), rnd.choice(LOCS), rnd.choice(["High", "Medium", "Low"]),
                rnd.choice(["Recruiter A", "Recruiter B"]), rnd.choice(["Manager X", "Manager Y", "Manager Z"]), status,
                opened, filled, rnd.choice(["Referral", "Job portal", "Agency", "Campus"]) if status == "Filled" else None,
                offer, rnd.randint(8, 60) * 1000 if status == "Filled" else None, stage])

# Training
trn = wb.create_sheet("Training")
trn.append(["Employee ID", "Training Program", "Training Category", "Delivery Mode", "Location", "Training Hours",
            "Training Date", "Completion Status", "Certification", "Pre Score", "Post Score", "Training Cost"])
progs = {"Safety Basics": "Compliance", "Data Ethics": "Compliance", "Excel Skills": "Technical",
         "Leadership 101": "Leadership", "Communication": "Soft skills"}
for _ in range(420):
    p = rnd.choice(list(progs))
    pre = rnd.randint(40, 70)
    trn.append([f"DEMO{rnd.randint(1, 160):03d}", p, progs[p], rnd.choice(["Classroom", "Online", "Blended"]),
                rnd.choice(LOCS), rnd.choice([2, 4, 8]), d(rnd.randint(1, 400)),
                rnd.choice(["Completed"] * 5 + ["Not completed"]), rnd.choice(["Yes", "No", None]),
                pre, pre + rnd.randint(-3, 25), rnd.randint(5, 40) * 100])

# Performance
pms = wb.create_sheet("PMS FY2025-26")
pms.append(["Employee ID", "Function", "Level", "Review Cycle", "Performance Rating", "Review Status", "Goals Set",
            "Goals Achieved", "IDP Created"])
for i in range(1, 141):
    done = rnd.random() < 0.85
    g = rnd.randint(3, 6)
    pms.append([f"DEMO{i:03d}", rnd.choice(DEPTS), rnd.choice(["L1", "L2", "L3"]), "FY2025-26",
                rnd.choice([2, 3, 3, 4, 4, 5]) if done else None, "Completed" if done else "Pending",
                g if done else None, rnd.randint(1, g) if done else None, rnd.choice(["Yes", "No"])])

# Employee relations
er = wb.create_sheet("Employee Relations")
er.append(["Case ID", "Issue Category", "Case Status", "Case Open Date", "Case Close Date", "Department", "Location"])
for i in range(1, 61):
    o = d(rnd.randint(5, 400))
    closed = rnd.random() < 0.75
    er.append([f"CASE{i:03d}", rnd.choice(["Workplace conduct", "Pay query", "Policy clarification", "Workload"]),
               "Closed" if closed else rnd.choice(["Open", "Pending"]), o,
               o + timedelta(days=rnd.randint(3, 40)) if closed else None, rnd.choice(DEPTS), rnd.choice(LOCS)])

# Engagement survey
sv = wb.create_sheet("Engagement Survey")
sv.append(["Department", "Survey Invited", "Survey Responded", "Engagement Score"])
for dep in DEPTS:
    for loc in LOCS:
        inv = rnd.randint(8, 25)
        sv.append([dep, inv, rnd.randint(int(inv * .6), inv), round(rnd.uniform(3.2, 4.5), 2)])

out = sys.argv[1] if len(sys.argv) > 1 else "SWITCH_demo_workbook_FICTIONAL.xlsx"
wb.save(out)
print("saved", out)
